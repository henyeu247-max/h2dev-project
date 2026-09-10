"""Extract spreadsheet evidence or timestamped video samples; never execute source tools."""
import argparse, json, pathlib, shutil, subprocess
from html.parser import HTMLParser
ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / '_audit/20260906-learning-corpus'

class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__();self.parts=[];self.links=[];self.hidden=0
    def handle_starttag(self,tag,attrs):
        if tag in ['script','style']:self.hidden+=1
        if tag in ['p','br','li','div']:self.parts.append('\n')
        if tag=='a':
            href=dict(attrs).get('href')
            if href:self.links.append(href)
    def handle_endtag(self,tag):
        if tag in ['script','style']:self.hidden=max(0,self.hidden-1)
    def handle_data(self,data):
        if not self.hidden:self.parts.append(data)

def descriptions():
    rows=[]
    for folder in sorted((ROOT/'docs').glob('VIDEO-*')):
        row={'sku':folder.name}
        for name in ['README.md','description.txt','DESCRIPTION_EMPTY.md']:
            p=folder/name
            if p.exists():row[name]=p.read_text(encoding='utf-8-sig')
        p=folder/'description.html'
        if p.exists():
            parser=VisibleHTML();parser.feed(p.read_text(encoding='utf-8-sig'))
            row['html_visible']=''.join(parser.parts);row['html_links']=parser.links
        rows.append(row)
    (OUT/'description-evidence.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf8')
    print('folders',len(rows),'html',sum('html_visible' in x for x in rows))

def sheets():
    import openpyxl
    result=[]
    for p in sorted(ROOT.rglob('*.xlsx')):
        if any(x in p.parts for x in ['node_modules','_backup','_archive','_audit']): continue
        w=openpyxl.load_workbook(p, data_only=False)
        for s in w:
            cells=[]
            for row in s:
                for c in row:
                    if c.value is not None or c.hyperlink:
                        cells.append({'cell':c.coordinate,'value':str(c.value) if c.value is not None else '', 'link':c.hyperlink.target if c.hyperlink else None})
            result.append({'path':str(p),'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'cells':cells})
    (OUT/'spreadsheet-evidence.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    for s in result:
        print(s['path'],s['sheet'],len(s['cells']))

def frames(sku, count):
    from PIL import Image, ImageDraw
    p=ROOT/'video'/sku/(sku+'.mp4')
    probe=json.loads(subprocess.check_output([shutil.which('ffprobe'),'-v','error','-show_format','-of','json',str(p)]))
    duration=float(probe['format']['duration'])
    folder=OUT/'visual-recovery'/sku;folder.mkdir(parents=True,exist_ok=True)
    records=[]; width=640; height=384
    sheet=Image.new('RGB',(width*3,height*((count+2)//3)),'white')
    draw=ImageDraw.Draw(sheet)
    for i in range(count):
        sec=duration*(i+.5)/count
        target=folder/f'{i:02d}-{sec:.1f}.jpg'
        subprocess.run([shutil.which('ffmpeg'),'-hide_banner','-loglevel','error','-ss',str(sec),'-i',str(p),'-frames:v','1','-vf','scale=640:-1','-y',str(target)],check=True)
        with Image.open(target) as im: sheet.paste(im,(i%3*width,i//3*height+24))
        draw.text((i%3*width+5,i//3*height+4),f'{sku} | {sec:.1f}s',fill='black')
        records.append({'seconds':sec,'path':str(target)})
    sheet.save(folder/'contact.jpg',quality=90)
    (folder/'manifest.json').write_text(json.dumps({'source':str(p),'duration':duration,'method':'uniform timestamp samples, not full video review','frames':records},indent=2),encoding='utf8')
    print(folder/'contact.jpg')

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--sheets',action='store_true');a.add_argument('--descriptions',action='store_true');a.add_argument('--sku');a.add_argument('--count',type=int,default=12);v=a.parse_args()
    if v.sheets:sheets()
    if v.descriptions:descriptions()
    if v.sku:frames(v.sku,v.count)
